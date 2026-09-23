import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
blue_header = "002060"
gray_border = "E2E8F0"
font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
fill_header = PatternFill(start_color=blue_header, end_color=blue_header, fill_type="solid")

ws_base = wb.active
ws_base.title = "BASE_CONTATOS"
headers = ["ID", "Concessão", "Área", "Função", "Nome", "E-mail", "Gestor", "Observação", "Status"]
ws_base.append(headers)

for col_num, h_text in enumerate(headers, 1):
    c = ws_base.cell(row=1, column=col_num)
    c.font = font_header
    c.fill = fill_header

wb.save("GCC_Pontos_Focais_e_Mailings_v4.xlsx")
print("Planilha GCC_Pontos_Focais_e_Mailings_v4.xlsx gerada com sucesso!")
